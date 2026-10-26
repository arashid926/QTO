import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from qto import boq, dxf_takeoff, samples, schedules
from qto import disciplines as dz

M = dz.MEASURE_COLUMNS


@pytest.fixture(scope="module")
def sample():
    return dxf_takeoff.takeoff(samples.sample_dxf_bytes())


def layer(result, name):
    return result.layers.set_index("Layer").loc[name]


def test_sample_units_detected(sample):
    assert sample.units_label == "Millimetres"
    assert sample.scale_to_m == 0.001


def test_wall_lengths(sample):
    # outer 42 m + inner 40.4 m perimeters, 2 x 8.6 m + 2 x 5.7 m internal wall faces
    assert layer(sample, "A-WALL")[M["Length"]] == pytest.approx(42 + 40.4 + 17.2 + 11.4, rel=1e-6)


def test_columns_and_slab(sample):
    cols = layer(sample, "S-COLUMN")
    assert cols[M["Count"]] == 6
    assert cols[M["Area"]] == pytest.approx(6 * 0.16)
    assert layer(sample, "S-SLAB")[M["Area"]] == pytest.approx(108.0)
    assert layer(sample, "S-COLUMN")["Measure"] == "Count"
    assert layer(sample, "S-SLAB")["Measure"] == "Area"


def test_floor_finish_area(sample):
    assert layer(sample, "A-FLOOR-FINISH")[M["Area"]] == pytest.approx(5.7 * 8.6 + 2 * 5.7 * 4.2)


def test_block_counts_and_classification(sample):
    blocks = sample.blocks.set_index("Block")
    assert blocks.loc["DOOR-900", "Count"] == 4
    assert blocks.loc["WIN-1200", "Count"] == 6
    assert blocks.loc["LIGHT-600", "Count"] == 8
    assert blocks.loc["DOOR-900", "Discipline"] == "Architectural"
    assert blocks.loc["LIGHT-600", "Discipline"] == "MEP"


def test_tag_counts(sample):
    counts = dxf_takeoff.count_tags(sample.texts, ["D1", "w1", "X9"]).set_index("Tag")["Count"]
    assert counts.to_dict() == {"D1": 4, "w1": 6, "X9": 0}


def test_text_only_layers_are_not_measured(sample):
    assert "A-TAG" not in set(sample.layers["Layer"])


def test_block_only_layer_not_billed(sample):
    assert not layer(sample, "A-DOOR")["Bill"]


def test_explode_blocks_adds_geometry():
    exploded = dxf_takeoff.takeoff(samples.sample_dxf_bytes(), explode_blocks=True)
    assert layer(exploded, "A-DOOR")[M["Length"]] > 0


def test_classify():
    assert dz.classify("S-BEAM") == ("Structural", "11 In-situ concrete works")
    assert dz.classify("E-LITE") == ("MEP", "39 Electrical services")
    assert dz.classify("C-ROAD-KERB")[0] == "Infrastructure"
    assert dz.classify("random")[1] == dz.UNCLASSIFIED
    assert dz.pretty_name("A-FLOOR-FINISH") == "Floor Finish"


def test_schedule_header_detection_and_mapping():
    csv = "Door schedule,,,\nProject X,,,\nMark,Description,Qty,Unit\nD1,Flush door 900,12,nos\nD2,Fire door,3,no.\n,Subtotal,,\n"
    sheets = schedules.read_schedule(csv.encode(), "doors.csv")
    table = sheets["CSV"]
    mapping = schedules.auto_map(table.columns)
    assert mapping["Ref"] == "Mark" and mapping["Quantity"] == "Qty" and mapping["Unit"] == "Unit"
    lines = schedules.normalise(table, mapping, "Architectural", "doors.csv")
    assert lines["Quantity"].tolist() == [12, 3]
    assert lines["Unit"].tolist() == ["nr", "nr"]
    assert lines["Work Section"].iloc[0] == "24 Doors, shutters and hatches"


def test_templates_round_trip():
    for discipline in dz.DISCIPLINES:
        t = schedules.template(discipline)
        table = schedules.read_schedule(t.to_csv(index=False).encode(), "t.csv")["CSV"]
        lines = schedules.normalise(table, schedules.auto_map(table.columns), discipline, "t")
        assert len(lines) == len(t)
        assert (lines["Rate"] > 0).all()


def test_boq_from_measured_and_export(sample):
    lines = boq.from_measured(sample.layers, "sample · layers")
    assert set(lines["Unit"]) <= {"m", "m²", "nr"}
    assert lines["Ref"].str.len().gt(0).all()
    priced = lines.assign(Rate=10.0)
    final = boq.finalise(priced)
    assert final["Amount"].sum() == pytest.approx(final["Quantity"].sum() * 10, rel=1e-6)
    assert boq.finalise(final).equals(final)  # idempotent, the BOQ editor relies on this

    workbook = load_workbook(io.BytesIO(boq.to_excel(final, {"Project": "Test"}, "USD", {"sample": sample.layers})))
    assert workbook.sheetnames[0] == "Summary"
    assert "Architectural" in workbook.sheetnames and "All Items" in workbook.sheetnames


def test_boq_handles_new_blank_rows():
    df = pd.DataFrame([{"Discipline": "Civil", "Description": "Excavation", "Quantity": "12.5", "Rate": None}])
    out = boq.finalise(df)
    assert out.loc[0, "Ref"] == "C.00.001"
    assert out.loc[0, "Amount"] == 0
