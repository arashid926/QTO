import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from qto import boq, measure


@pytest.fixture
def state():
    s = measure.new_state()
    s["calibration"][measure.page_key("plan.pdf", 0)] = 0.01  # 1 drawing unit = 10 mm
    return s


def test_units_follow_dimensions():
    assert measure.item_unit("Count", 0, 0) == "nr"
    assert measure.item_unit("Length", 0, 0) == "m"
    assert measure.item_unit("Length", 3.0, 0) == "m²"
    assert measure.item_unit("Length", 0.6, 0.4) == "m³"
    assert measure.item_unit("Area", 0, 0) == "m²"
    assert measure.item_unit("Area", 0, 0.2) == "m³"


def test_wall_area_with_deduction(state):
    wall = measure.add_item(state, "Blockwork wall", "Architectural", "14 Masonry", "Length", width=3.0)
    measure.add_shape(state, wall["id"], "plan.pdf", 0, [[0, 0], [1000, 0], [1000, 500]])  # 15 m run
    door = measure.add_shape(state, wall["id"], "plan.pdf", 0, [[0, 0], [90, 0]])  # 0.9 m opening
    door.update({"deduct": True, "timesing": 2})
    sheet = measure.qto_sheet(state)
    assert sheet["Quantity"].tolist() == [pytest.approx(45.0), pytest.approx(-5.4)]
    line = measure.boq_lines(state).iloc[0]
    assert (line["Ref"], line["Unit"], line["Quantity"]) == ("TO-001", "m²", pytest.approx(39.6))


def test_slab_volume_and_count(state):
    slab = measure.add_item(state, "RC slab", "Structural", "11 In-situ concrete works", "Area", depth=0.2)
    measure.add_shape(state, slab["id"], "plan.pdf", 0, [[0, 0], [1000, 0], [1000, 800], [0, 800]])  # 10 x 8 m
    doors = measure.add_item(state, "Doors", "Architectural", "24 Doors, shutters and hatches", "Count")
    measure.add_shape(state, doors["id"], "other.pdf", 3, [[1, 1], [2, 2], [3, 3]])  # counts need no scale
    lines = measure.boq_lines(state).set_index("Ref")
    assert lines.loc["TO-001", "Quantity"] == pytest.approx(16.0) and lines.loc["TO-001", "Unit"] == "m³"
    assert lines.loc["TO-002", "Quantity"] == 3 and lines.loc["TO-002", "Unit"] == "nr"


def test_uncalibrated_page_has_no_quantity():
    s = measure.new_state()
    item = measure.add_item(s, "Skirting", "Architectural", "28 Floor, wall, ceiling and roof finishings", "Length")
    measure.add_shape(s, item["id"], "scan.png", 0, [[0, 0], [100, 0]])
    assert measure.qto_sheet(s)["Quantity"].isna().all()
    assert measure.boq_lines(s).empty


def test_scale_and_calibration():
    assert measure.metres_per_unit_from_scale(100) * 72 == pytest.approx(2.54)  # one inch on paper at 1:100
    assert measure.calibrate([[0, 0], [300, 400]], 5.0) == pytest.approx(0.01)


def test_sheet_edits_and_deletes(state):
    item = measure.add_item(state, "Skirting", "Architectural", "28 Floor, wall, ceiling and roof finishings", "Length")
    a = measure.add_shape(state, item["id"], "plan.pdf", 0, [[0, 0], [100, 0]])
    measure.add_shape(state, item["id"], "plan.pdf", 0, [[0, 0], [200, 0]])
    sheet = measure.qto_sheet(state)
    sheet.loc[0, "Timesing"] = 4
    assert measure.update_from_sheet(state, sheet.iloc[[0]])
    assert [s["id"] for s in state["shapes"]] == [a["id"]] and state["shapes"][0]["timesing"] == 4
    assert not measure.update_from_sheet(state, measure.qto_sheet(state))  # stable: no change, no rerun loop


def test_json_round_trip(state):
    item = measure.add_item(state, "Floor tiles", "Architectural", "28 Floor, wall, ceiling and roof finishings", "Area")
    measure.add_shape(state, item["id"], "plan.pdf", 0, [[0, 0], [100, 0], [100, 100]])
    restored = measure.from_json(measure.to_json(state))
    assert measure.qto_sheet(restored).drop(columns="Shape").equals(measure.qto_sheet(state).drop(columns="Shape"))
    with pytest.raises(ValueError):
        measure.from_json(b'{"format": "something else"}')


def test_excel_links_boq_to_qto_sheet(state):
    wall = measure.add_item(state, "Blockwork wall", "Architectural", "14 Masonry", "Length", width=3.0)
    measure.add_shape(state, wall["id"], "plan.pdf", 0, [[0, 0], [1000, 0]])
    lines = boq.finalise(measure.boq_lines(state))
    wb = load_workbook(io.BytesIO(boq.to_excel(lines, qto=measure.qto_sheet(state))))
    assert boq.QTO_SHEET in wb.sheetnames
    bill = wb["Architectural"]
    quantity_cells = [bill.cell(row=r, column=4).value for r in range(4, bill.max_row + 1)]
    assert any(isinstance(v, str) and v.startswith("=SUMIF('QTO Sheet'") for v in quantity_cells)
    qto = wb[boq.QTO_SHEET]
    assert qto["A5"].value == "TO-001" and str(qto["M5"].value).startswith("=IF(H5=")


def test_from_measured_factor_and_height():
    table = pd.DataFrame([
        {"Bill": True, "Discipline": "Architectural", "Work Section": "14 Masonry", "Measure": "Length",
         "Description": "Walls (double line)", "Factor": 0.5, "× Height/Thk (m)": 3.0, "Length (m)": 40.0},
        {"Bill": True, "Discipline": "Structural", "Work Section": "11 In-situ concrete works", "Measure": "Area",
         "Description": "Slab", "Factor": 1.0, "× Height/Thk (m)": 0.2, "Area (m²)": 100.0},
    ])
    lines = boq.from_measured(table, "test")
    assert lines["Quantity"].tolist() == [pytest.approx(60.0), pytest.approx(20.0)]
    assert lines["Unit"].tolist() == ["m²", "m³"]
