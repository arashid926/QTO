"""Bill of Quantities: build lines from take-offs and export a priced Excel workbook."""

from __future__ import annotations

import io
import re
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import disciplines as dz

COLUMNS = ["Ref", "Discipline", "Work Section", "Description", "Unit", "Quantity", "Rate", "Amount", "Source"]
TEXT_COLUMNS = ["Ref", "Discipline", "Work Section", "Description", "Unit", "Source"]


def empty() -> pd.DataFrame:
    df = pd.DataFrame(columns=COLUMNS)
    df[TEXT_COLUMNS] = df[TEXT_COLUMNS].astype(str)
    return df.astype({"Quantity": float, "Rate": float, "Amount": float})


def from_measured(table: pd.DataFrame, source: str) -> pd.DataFrame:
    """Turn ticked rows of a take-off table into BOQ lines using each row's chosen measure."""
    rows = []
    for _, r in table[table["Bill"].fillna(False).astype(bool)].iterrows():
        measure = r.get("Measure") or "Count"
        qty = pd.to_numeric(r.get(dz.MEASURE_COLUMNS.get(measure, "Count")), errors="coerce")
        if pd.isna(qty) or qty <= 0:
            continue
        rows.append({
            "Ref": "",
            "Discipline": r.get("Discipline") or dz.DEFAULT_CLASS[0],
            "Work Section": r.get("Work Section") or dz.UNCLASSIFIED,
            "Description": str(r.get("Description") or "").strip() or "Item",
            "Unit": dz.MEASURE_UNITS.get(measure, "nr"),
            "Quantity": round(float(qty), 3),
            "Rate": 0.0,
            "Source": source,
        })
    return finalise(pd.DataFrame(rows, columns=COLUMNS), sort=False)


def _discipline_order(name: str) -> int:
    return dz.DISCIPLINES.index(name) if name in dz.DISCIPLINES else len(dz.DISCIPLINES)


def finalise(df: pd.DataFrame, sort: bool = True) -> pd.DataFrame:
    """Fill missing columns, recompute amounts and number blank refs. Safe to call repeatedly."""
    out = df.copy()
    for column in COLUMNS:
        if column not in out:
            out[column] = "" if column in TEXT_COLUMNS else 0.0
    out = out[COLUMNS]
    for column in TEXT_COLUMNS:
        out[column] = out[column].fillna("").astype(str)
    out["Quantity"] = pd.to_numeric(out["Quantity"], errors="coerce").fillna(0.0).astype(float)
    out["Rate"] = pd.to_numeric(out["Rate"], errors="coerce").fillna(0.0).astype(float)
    out["Amount"] = (out["Quantity"] * out["Rate"]).round(2)
    if sort and not out.empty:
        out = out.assign(_o=out["Discipline"].map(_discipline_order)).sort_values(["_o", "Work Section"], kind="stable").drop(columns="_o")
    out = out.reset_index(drop=True)
    counters: dict[tuple[str, str], int] = {}
    for i in out.index:
        key = (dz.DISCIPLINE_CODES.get(out.at[i, "Discipline"], "X"), dz.section_number(out.at[i, "Work Section"]))
        counters[key] = counters.get(key, 0) + 1
        if not out.at[i, "Ref"].strip():
            out.at[i, "Ref"] = f"{key[0]}.{key[1]}.{counters[key]:03d}"
    return out


def summary(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["Discipline", "Items", "Amount"])
    g = df.groupby("Discipline").agg(Items=("Description", "size"), Amount=("Amount", "sum")).reset_index()
    return g.sort_values("Discipline", key=lambda s: s.map(_discipline_order)).reset_index(drop=True)


# ---------------------------------------------------------------- Excel export

NAVY, BLUE, PALE, GREY = "0F172A", "1E3A8A", "E0E7FF", "F1F5F9"
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
QTY_FMT, MONEY_FMT = "#,##0.000", "#,##0.00"


def _sheet_name(name: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in {t.lower() for t in taken}:
        suffix = f" ({n})"
        candidate, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(candidate)
    return candidate


def _header(ws, row: int, values: list[str]) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _write_discipline(ws, discipline: str, lines: pd.DataFrame, currency: str) -> str:
    """Write one discipline bill; returns the address of its total cell."""
    ws["A1"] = f"{discipline} — Bill of Quantities"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    _header(ws, 3, ["Ref", "Description", "Unit", "Quantity", f"Rate ({currency})", f"Amount ({currency})", "Source"])
    row, subtotals = 4, []
    for section, group in lines.groupby("Work Section", sort=False):
        ws.cell(row=row, column=2, value=section).font = Font(bold=True, color=BLUE)
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=PALE)
        row += 1
        first = row
        for _, line in group.iterrows():
            values = [line["Ref"], line["Description"], line["Unit"], float(line["Quantity"]), float(line["Rate"]), f"=D{row}*E{row}", line["Source"]]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=col == 2)
            ws.cell(row=row, column=4).number_format = QTY_FMT
            ws.cell(row=row, column=5).number_format = MONEY_FMT
            ws.cell(row=row, column=6).number_format = MONEY_FMT
            row += 1
        ws.cell(row=row, column=2, value=f"Total — {section}").font = Font(bold=True)
        total = ws.cell(row=row, column=6, value=f"=SUM(F{first}:F{row - 1})")
        total.font, total.number_format = Font(bold=True), MONEY_FMT
        subtotals.append(f"F{row}")
        row += 2
    ws.cell(row=row, column=2, value=f"TOTAL {discipline.upper()} CARRIED TO SUMMARY").font = Font(bold=True, color="FFFFFF")
    grand = ws.cell(row=row, column=6, value="=" + "+".join(subtotals) if subtotals else 0)
    grand.font, grand.number_format = Font(bold=True, color="FFFFFF"), MONEY_FMT
    for col in range(1, 8):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    for col, width in zip("ABCDEFG", (12, 62, 8, 14, 14, 18, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "3:3"
    return f"F{row}"


def _write_table(ws, df: pd.DataFrame) -> None:
    _header(ws, 1, [str(c) for c in df.columns])
    for r, values in enumerate(df.itertuples(index=False), start=2):
        for c, value in enumerate(values, start=1):
            if isinstance(value, float) and pd.isna(value):
                value = None
            elif not isinstance(value, (int, float, str, bool)) and value is not None:
                value = str(value)
            ws.cell(row=r, column=c, value=value)
    for c, column in enumerate(df.columns, start=1):
        width = max([len(str(column))] + [len(str(v)) for v in df[column].head(200)]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 8), 60)
    ws.freeze_panes = "A2"


def to_excel(
    df: pd.DataFrame,
    project: dict[str, str] | None = None,
    currency: str = "",
    details: dict[str, pd.DataFrame] | None = None,
) -> bytes:
    """Summary sheet + one priced bill per discipline (live formulas) + raw take-off sheets."""
    df = finalise(df)
    project = project or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    taken = {"Summary"}

    ws["A1"] = "BILL OF QUANTITIES"
    ws["A1"].font = Font(bold=True, size=18, color=NAVY)
    ws["A2"] = "Measured generally in accordance with RICS NRM2"
    ws["A2"].font = Font(italic=True, color="64748B")
    row = 4
    for label, value in [*project.items(), ("Date", date.today().strftime("%d %b %Y"))]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    _header(ws, row, ["No.", "Discipline", "Items", f"Amount ({currency})"])
    first = row + 1
    disciplines = sorted(df["Discipline"].unique(), key=_discipline_order)
    for n, discipline in enumerate(disciplines, start=1):
        lines = df[df["Discipline"] == discipline]
        sheet = _sheet_name(discipline or "General", taken)
        total_cell = _write_discipline(wb.create_sheet(sheet), discipline, lines, currency)
        row += 1
        for col, value in enumerate([n, discipline, len(lines), f"='{sheet}'!{total_cell}"], start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = BORDER
        ws.cell(row=row, column=4).number_format = MONEY_FMT
    row += 1
    ws.cell(row=row, column=2, value="GRAND TOTAL").font = Font(bold=True, color="FFFFFF")
    grand = ws.cell(row=row, column=4, value=f"=SUM(D{first}:D{row - 1})" if disciplines else 0)
    grand.font, grand.number_format = Font(bold=True, color="FFFFFF"), MONEY_FMT
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    for col, width in zip("ABCD", (22, 40, 10, 22)):
        ws.column_dimensions[col].width = width

    _write_table(wb.create_sheet(_sheet_name("All Items", taken)), df)
    for name, table in (details or {}).items():
        if isinstance(table, pd.DataFrame) and not table.empty:
            _write_table(wb.create_sheet(_sheet_name(f"TO {name}", taken)), table.head(100_000))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
